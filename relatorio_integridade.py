# -*- coding: utf-8 -*-
"""
================================================================================
 RELATÓRIO DE INTEGRIDADE CONTÁBIL E ORÇAMENTÁRIA — Governo do Distrito Federal
================================================================================

Lê os 6 demonstrativos OFICIAIS publicados em
"13 - DEMONSTRATIVOS CONTÁBEIS/{ano}/{subpasta}/Lista{Tipo} {mes:02d}.pdf"
e aplica as regras de integridade cruzando-os entre si (leitor_pdf_oficial.py
+ motor_regras.py), gerando Relatorio_Integridade.xlsx e .pdf.

REESCRITO em 18/08/2026: a versão anterior dependia de 5 módulos
(leitor_dados_mensais.py, motor_regras.py, gerar_excel.py, gerar_pdf.py,
catalogo_campos.py) que não existiam mais no projeto e nunca haviam sido
commitados no git -- foram perdidos. Esta versão elimina também o passo de
digitação manual em dados_mensais.xlsx: lê os PDFs oficiais diretamente,
o que permite rodar sozinho na rotina noturna sem intervenção humana.

Uso:
    python relatorio_integridade.py --mes 7 --ano 2026
    python relatorio_integridade.py --mes 7 --ano 2026 --formato pdf
================================================================================
"""
import argparse, sys
from datetime import datetime
from pathlib import Path

if sys.platform == "win32":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.platypus import (SimpleDocTemplate, Table, TableStyle,
                                    Paragraph, Spacer)
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT
    REPORTLAB_OK = True
except ImportError:
    REPORTLAB_OK = False

from leitor_pdf_oficial import (carregar, PdfNaoEncontrado, pasta13_mudou,
                                marcar_pasta13_processada, ultimo_mes_disponivel)
from motor_regras import run_all_rules, Achado

OUTPUT_DIR = Path(__file__).parent
MESES = {1:'Janeiro',2:'Fevereiro',3:'Março',4:'Abril',5:'Maio',
         6:'Junho',7:'Julho',8:'Agosto',9:'Setembro',
         10:'Outubro',11:'Novembro',12:'Dezembro'}

COR_STATUS = {
    'OK':          ('1F5C2E', 'D9F0DD'),
    'DIVERGENTE':  ('B00000', 'FAD7DA'),
    'ALERTA':      ('9C6500', 'FCF0CE'),
    'INFORMATIVO': ('1F4E78', 'DDEBF7'),
    'NAO_VERIF':   ('595959', 'E7E6E6'),
}


def _hx(h):
    """Hex sem '#' (usado no Excel/openpyxl) -> Color do reportlab (que
    exige o '#')."""
    return colors.HexColor('#' + h)
STATUS_LABEL = {'NAO_VERIF': 'NÃO VERIF.'}


# ─────────────────────────────────────────────────────────────────────────────
#  EXCEL
# ─────────────────────────────────────────────────────────────────────────────
def _fill(h): return PatternFill('solid', fgColor=h)
def _side(): return Side(border_style='thin', color='B8CCE4')
B_THIN = Border(left=_side(), right=_side(), top=_side(), bottom=_side())


def _cel(ws, r, c, v='', bold=False, sz=10, bg=None, fg=None, ha='left', wrap=False):
    x = ws.cell(row=r, column=c, value=v)
    x.font = Font(name='Arial', bold=bold, size=sz, color=fg or '000000')
    x.alignment = Alignment(horizontal=ha, vertical='center', wrap_text=wrap)
    if bg: x.fill = _fill(bg)
    x.border = B_THIN
    return x


def gerar_excel(achados, mes, ano, output_path):
    wb = openpyxl.Workbook(); ws = wb.active
    ws.title = 'Relatório de Integridade'
    ws.sheet_view.showGridLines = False
    widths = {'A': 8, 'B': 46, 'C': 30, 'D': 12, 'E': 18, 'F': 30,
              'G': 18, 'H': 30, 'I': 16, 'J': 55}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    ws.merge_cells('A1:J2')
    c = ws['A1']
    c.value = (f'GOVERNO DO DISTRITO FEDERAL\n'
               f'Relatório de Integridade Contábil e Orçamentária — '
               f'{MESES[mes]}/{ano}')
    c.font = Font(name='Arial', bold=True, size=13, color='2E5C8A')
    c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 40

    header = ['Regra', 'Título', 'Base Normativa', 'Status', 'Valor A',
              'Rótulo A', 'Valor B', 'Rótulo B', 'Diferença', 'Observação']
    H = 4
    for j, htxt in enumerate(header, start=1):
        _cel(ws, H, j, htxt, bold=True, bg='2E5C8A', fg='FFFFFF', ha='center')

    for i, a in enumerate(achados, start=H + 1):
        _, bg = COR_STATUS.get(a.status, ('000000', 'FFFFFF'))
        status_txt = STATUS_LABEL.get(a.status, a.status)
        _cel(ws, i, 1, a.regra, bg=bg, ha='center')
        _cel(ws, i, 2, a.titulo, bg=bg, wrap=True)
        _cel(ws, i, 3, a.base_normativa, bg=bg, wrap=True)
        _cel(ws, i, 4, status_txt, bold=True, bg=bg, ha='center')
        _cel(ws, i, 5, a.valor_a, bg=bg, ha='right')
        if a.valor_a is not None: ws.cell(row=i, column=5).number_format = '#,##0.00'
        _cel(ws, i, 6, a.rotulo_a, bg=bg, wrap=True)
        _cel(ws, i, 7, a.valor_b, bg=bg, ha='right')
        if a.valor_b is not None: ws.cell(row=i, column=7).number_format = '#,##0.00'
        _cel(ws, i, 8, a.rotulo_b, bg=bg, wrap=True)
        dif = a.diferenca
        _cel(ws, i, 9, dif, bg=bg, ha='right')
        if dif is not None: ws.cell(row=i, column=9).number_format = '#,##0.00'
        _cel(ws, i, 10, a.observacao, bg=bg, wrap=True)

    ws.freeze_panes = f'A{H+1}'
    wb.save(output_path)
    print(f"  Excel salvo: {output_path}")


# ─────────────────────────────────────────────────────────────────────────────
#  PDF
# ─────────────────────────────────────────────────────────────────────────────
def _brl(v):
    if v is None:
        return '—'
    return f'{v:,.2f}'.replace(',', '§').replace('.', ',').replace('§', '.')


def gerar_pdf(achados, mes, ano, output_path):
    if not REPORTLAB_OK:
        print("  AVISO: reportlab não instalado -- PDF não gerado.")
        return

    doc = SimpleDocTemplate(str(output_path), pagesize=A4,
                             leftMargin=12*mm, rightMargin=12*mm,
                             topMargin=12*mm, bottomMargin=12*mm)
    styles = getSampleStyleSheet()
    elems = []

    st_t = ParagraphStyle('t', parent=styles['Normal'], fontName='Helvetica-Bold',
                           fontSize=15, textColor=colors.HexColor('#2E5C8A'), leading=18)
    st_s = ParagraphStyle('s', parent=styles['Normal'], fontSize=9, textColor=colors.grey)
    elems.append(Paragraph('GOVERNO DO DISTRITO FEDERAL', st_s))
    elems.append(Paragraph(f'Relatório de Integridade Contábil e Orçamentária — '
                            f'{MESES[mes]}/{ano}', st_t))
    elems.append(Paragraph(f'Gerado em {datetime.now():%d/%m/%Y às %H:%M:%S} · '
                            f'Fonte: PDFs oficiais publicados (PSIAG550) — '
                            f'BF, BO, BP, DFC, DVP, DMPL', st_s))
    elems.append(Spacer(1, 6*mm))

    contagem = {}
    for a in achados:
        contagem[a.status] = contagem.get(a.status, 0) + 1
    ordem = ['OK', 'DIVERGENTE', 'ALERTA', 'INFORMATIVO', 'NAO_VERIF']
    painel_dados = [[str(contagem.get(s, 0)) for s in ordem],
                     [STATUS_LABEL.get(s, s) for s in ordem]]
    painel = Table([[str(len(achados))] + painel_dados[0],
                    ['Total'] + painel_dados[1]],
                   colWidths=[28*mm]*6)
    ts_painel = [('GRID', (0, 0), (-1, -1), 0.4, colors.white),
                 ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                 ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                 ('FONTSIZE', (0, 0), (-1, 0), 16),
                 ('FONTSIZE', (0, 1), (-1, 1), 8),
                 ('TOPPADDING', (0, 0), (-1, -1), 8),
                 ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
                 ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#2E5C8A')),
                 ('TEXTCOLOR', (0, 0), (0, -1), colors.white)]
    for i, s in enumerate(ordem, start=1):
        fg, bg = COR_STATUS[s]
        ts_painel.append(('BACKGROUND', (i, 0), (i, -1), _hx(bg)))
        ts_painel.append(('TEXTCOLOR', (i, 0), (i, 0), _hx(fg)))
    painel.setStyle(TableStyle(ts_painel))
    elems.append(painel)
    elems.append(Spacer(1, 6*mm))

    prioritarios = [a for a in achados if a.status in ('DIVERGENTE', 'ALERTA')]
    if prioritarios:
        st_h2 = ParagraphStyle('h2', parent=styles['Normal'], fontName='Helvetica-Bold',
                                fontSize=12, textColor=colors.HexColor('#2E5C8A'))
        elems.append(Paragraph('ACHADOS PRIORITÁRIOS (DIVERGENTE / ALERTA)', st_h2))
        elems.append(Spacer(1, 2*mm))
        st_r = ParagraphStyle('r', fontName='Helvetica-Bold', fontSize=9)
        st_n = ParagraphStyle('n', fontName='Helvetica', fontSize=8.5, leading=11)
        st_o = ParagraphStyle('o', fontName='Helvetica-Oblique', fontSize=8, leading=10,
                              textColor=colors.HexColor('#444444'))
        for a in prioritarios:
            fg, bg = COR_STATUS[a.status]
            linhas = [[Paragraph(f'Regra {a.regra} — {a.titulo}', st_r),
                      Paragraph(a.status, st_r)]]
            tab = Table(linhas, colWidths=[150*mm, 26*mm])
            tab.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, -1), _hx(bg)),
                ('TEXTCOLOR', (1, 0), (1, 0), _hx(fg)),
                ('ALIGN', (1, 0), (1, 0), 'CENTER'),
                ('TOPPADDING', (0, 0), (-1, -1), 4), ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ]))
            elems.append(tab)
            corpo = [[Paragraph(f'{a.rotulo_a or "Valor A"}: {_brl(a.valor_a)}', st_n),
                     Paragraph(f'{a.rotulo_b or "Valor B"}: {_brl(a.valor_b)}', st_n)]]
            if a.diferenca is not None:
                corpo.append([Paragraph(f'Diferença apurada: {_brl(a.diferenca)}', st_r), ''])
            tab2 = Table(corpo, colWidths=[88*mm, 88*mm])
            tab2.setStyle(TableStyle([('GRID', (0, 0), (-1, -1), 0.3, colors.lightgrey),
                                      ('TOPPADDING', (0, 0), (-1, -1), 3),
                                      ('BOTTOMPADDING', (0, 0), (-1, -1), 3)]))
            elems.append(tab2)
            if a.base_normativa or a.observacao:
                txt = (f'<b>Base normativa:</b> {a.base_normativa}. ' if a.base_normativa else '')
                txt += a.observacao
                elems.append(Paragraph(txt, st_o))
            elems.append(Spacer(1, 3*mm))
        elems.append(Spacer(1, 4*mm))

    st_h2 = ParagraphStyle('h2b', parent=styles['Normal'], fontName='Helvetica-Bold',
                            fontSize=12, textColor=colors.HexColor('#2E5C8A'))
    elems.append(Paragraph('DETALHAMENTO COMPLETO', st_h2))
    elems.append(Spacer(1, 2*mm))
    st_cell = ParagraphStyle('c', fontName='Helvetica', fontSize=7, leading=8.5)
    st_cell_b = ParagraphStyle('cb', fontName='Helvetica-Bold', fontSize=7, leading=8.5)
    st_val = ParagraphStyle('v', fontName='Helvetica', fontSize=7, alignment=TA_RIGHT, leading=8.5)
    dados = [[Paragraph('Regra', st_cell_b), Paragraph('Título', st_cell_b),
              Paragraph('Valor A', st_cell_b), Paragraph('Valor B', st_cell_b),
              Paragraph('Diferença', st_cell_b), Paragraph('Status', st_cell_b)]]
    for a in achados:
        dados.append([Paragraph(str(a.regra), st_cell), Paragraph(a.titulo, st_cell),
                      Paragraph(_brl(a.valor_a), st_val), Paragraph(_brl(a.valor_b), st_val),
                      Paragraph(_brl(a.diferenca), st_val),
                      Paragraph(STATUS_LABEL.get(a.status, a.status), st_cell_b)])
    tab = Table(dados, colWidths=[14*mm, 70*mm, 30*mm, 30*mm, 25*mm, 22*mm], repeatRows=1)
    ts = [('GRID', (0, 0), (-1, -1), 0.3, colors.grey),
          ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
          ('TOPPADDING', (0, 0), (-1, -1), 2), ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
          ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2E5C8A')),
          ('TEXTCOLOR', (0, 0), (-1, 0), colors.white)]
    for i, a in enumerate(achados, start=1):
        _, bg = COR_STATUS.get(a.status, ('000000', 'FFFFFF'))
        ts.append(('BACKGROUND', (0, i), (-1, i), _hx(bg)))
    tab.setStyle(TableStyle(ts))
    elems.append(tab)

    doc.build(elems)
    print(f"  PDF salvo:   {output_path}")


# ─────────────────────────────────────────────────────────────────────────────
#  MAIN
# ─────────────────────────────────────────────────────────────────────────────
def main():
    p = argparse.ArgumentParser(
        description='Relatório de Integridade Contábil e Orçamentária - GDF '
                    '(lê os PDFs oficiais da pasta 13, não precisa de Oracle)')
    p.add_argument('--mes', type=str, required=True,
                   help="1-12, ou 'ultimo' para o último mês com os 6 "
                        "demonstrativos já publicados na pasta 13")
    p.add_argument('--ano', type=int, required=True)
    p.add_argument('--formato', choices=['ambos', 'excel', 'pdf'], default='ambos')
    p.add_argument('--somente-se-mudou', action='store_true',
                   help='Só roda (e só grava arquivos) se a pasta 13 do ano '
                        'pedido mudou desde a última vez que este script '
                        'gerou o relatório com sucesso. Usado na rotina '
                        'noturna para não subir Excel/PDF novos sem motivo.')
    a = p.parse_args()

    if a.mes == 'ultimo':
        mes_resolvido = ultimo_mes_disponivel(a.ano)
        if mes_resolvido is None:
            print(f"ERRO: nenhum mês de {a.ano} tem os 6 demonstrativos "
                  f"publicados na pasta 13.")
            sys.exit(1)
        a.mes = mes_resolvido
        print(f"  --mes ultimo -> resolvido para {a.mes:02d}/{a.ano}")
    else:
        try:
            a.mes = int(a.mes)
        except ValueError:
            print("ERRO: --mes deve ser um número de 1 a 12, ou 'ultimo'."); sys.exit(1)

    if not 1 <= a.mes <= 12:
        print("ERRO: --mes deve estar entre 1 e 12."); sys.exit(1)

    if a.somente_se_mudou and not pasta13_mudou(a.ano):
        print(f"  Pasta 13 ({a.ano}) sem alterações desde a última rodada — "
              f"nada a gerar. Use sem --somente-se-mudou para forçar.")
        sys.exit(0)

    print(f"\n{'='*70}")
    print(f"  RELATÓRIO DE INTEGRIDADE CONTÁBIL E ORÇAMENTÁRIA")
    print(f"  {MESES[a.mes]}/{a.ano}  |  Fonte: PDFs oficiais (pasta 13)")
    print(f"{'='*70}")

    print("\n[1/3] Localizando e lendo os 6 PDFs oficiais...")
    try:
        d = carregar(a.mes, a.ano)
    except PdfNaoEncontrado as e:
        print(f"ERRO: {e}")
        sys.exit(1)
    for tipo, caminho in d.caminhos.items():
        print(f"  {tipo:<5} -> {caminho}")

    print("\n[2/3] Aplicando as regras de integridade...")
    achados = run_all_rules(d)
    from collections import Counter
    contagem = Counter(x.status for x in achados)
    print(f"  Total de verificações: {len(achados)}")
    for status, n in contagem.most_common():
        print(f"    {STATUS_LABEL.get(status, status)}: {n}")

    divergentes = [x for x in achados if x.status == 'DIVERGENTE']
    if divergentes:
        print(f"\n  {len(divergentes)} DIVERGÊNCIA(S):")
        for x in divergentes:
            print(f"    Regra {x.regra}: {x.titulo}")
            print(f"      {x.rotulo_a}: {x.valor_a:,.2f}  |  {x.rotulo_b}: {x.valor_b:,.2f}"
                  f"  |  dif: {x.diferenca:,.2f}")

    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    base = OUTPUT_DIR / f'Relatorio_Integridade_{a.ano}_{a.mes:02d}_{timestamp}'
    print("\n[3/3] Gerando arquivos...")
    if a.formato in ('ambos', 'excel'):
        gerar_excel(achados, a.mes, a.ano, base.with_suffix('.xlsx'))
    if a.formato in ('ambos', 'pdf'):
        gerar_pdf(achados, a.mes, a.ano, base.with_suffix('.pdf'))

    if a.somente_se_mudou:
        marcar_pasta13_processada(a.ano)

    print(f"\n  Concluído em {datetime.now():%H:%M:%S}\n")
    sys.exit(1 if divergentes else 0)


if __name__ == "__main__":
    main()
