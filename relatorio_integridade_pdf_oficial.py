# -*- coding: utf-8 -*-
"""
================================================================================
 RELATÓRIO DE INTEGRIDADE CONTÁBIL E ORÇAMENTÁRIA — Governo do Distrito Federal
================================================================================

RENOMEADO em 10/09/2026 (era relatorio_integridade.py) para deixar explícito
que este script é o ÚNICO dos três diagnósticos do projeto que NÃO toca no
Oracle nem lê os .xlsx gerados por mestre.py/bo.py/bp.py/etc. — ele lê
SÓ os 6 demonstrativos OFICIAIS já PUBLICADOS EM PDF na pasta
"13 - DEMONSTRATIVOS CONTÁBEIS/{ano}/{subpasta}/Lista{Tipo} {mes:02d}.pdf"
e aplica as regras de integridade cruzando-os entre si (leitor_pdf_oficial.py
+ motor_regras.py), gerando Relatorio_Integridade.xlsx/.pdf + o painel
painel/painel_integridade_pdf_oficial.html. Os outros dois diagnósticos do
projeto usam fontes diferentes: diag.py/diag_tipoagreg.py consultam o Oracle
diretamente; auditoria_consolidada.py lê os .xlsx que mestre.py/bo.py/bp.py/
dfc.py/dmpl.py/dvp.py/gerar_balancete.py já geraram (sem Oracle e sem PDF).

REESCRITO em 18/08/2026: a versão anterior dependia de 5 módulos
(leitor_dados_mensais.py, motor_regras.py, gerar_excel.py, gerar_pdf.py,
catalogo_campos.py) que não existiam mais no projeto e nunca haviam sido
commitados no git -- foram perdidos. Esta versão elimina também o passo de
digitação manual em dados_mensais.xlsx: lê os PDFs oficiais diretamente,
o que permite rodar sozinho na rotina noturna sem intervenção humana.

Uso:
    python relatorio_integridade_pdf_oficial.py --mes 7 --ano 2026
    python relatorio_integridade_pdf_oficial.py --mes 7 --ano 2026 --formato pdf
================================================================================
"""
import argparse, json, sys
from datetime import datetime
from pathlib import Path

if sys.platform == "win32":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.platypus import (SimpleDocTemplate, Table, TableStyle,
                                    Paragraph, Spacer)
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT
    REPORTLAB_OK = True
except ImportError:
    REPORTLAB_OK = False

from leitor_pdf_oficial import (carregar, PdfNaoEncontrado, pasta13_mudou,
                                marcar_pasta13_processada, ultimo_mes_disponivel)
from motor_regras import run_all_rules, Achado

from saida.xlsx_html import fmt_brl
from auditoria_consolidada import CSS as _CSS_BASE  # reaproveita o mesmo tema visual

OUTPUT_DIR = Path(__file__).parent
MESES = {1:'Janeiro',2:'Fevereiro',3:'Março',4:'Abril',5:'Maio',
         6:'Junho',7:'Julho',8:'Agosto',9:'Setembro',
         10:'Outubro',11:'Novembro',12:'Dezembro'}

COR_STATUS = {
    'OK':          ('1F5C2E', 'D9F0DD'),
    'DIVERGENTE':  ('B00000', 'FAD7DA'),
    'ALERTA':      ('9C6500', 'FCF0CE'),
    'INFORMATIVO': ('1F4E78', 'DDEBF7'),
    'NAO_VERIF':   ('595959', 'E7E6E6'),
}

# ── Painel HTML: status do motor_regras.py (5 categorias) -> chip/ícone.
#    DIVERGENTE/NAO_VERIF não existem no vocabulário OK/ERRO/ALERTA/INFO
#    usado pelos outros painéis (saida/xlsx_html.py) -- mapeados à parte,
#    com c-gray novo (CSS_EXTRA) para NAO_VERIF.
ICONES_RI = {'OK': '✔', 'DIVERGENTE': '✘', 'ALERTA': '⚠',
             'INFORMATIVO': 'ℹ', 'NAO_VERIF': '•'}
CHIP_CLASS_RI = {'OK': 'c-ok', 'DIVERGENTE': 'c-err', 'ALERTA': 'c-alr',
                  'INFORMATIVO': 'c-inf', 'NAO_VERIF': 'c-gray'}


def _hx(h):
    """Hex sem '#' (usado no Excel/openpyxl) -> Color do reportlab (que
    exige o '#')."""
    return colors.HexColor('#' + h)
STATUS_LABEL = {'NAO_VERIF': 'NÃO VERIF.'}


# ─────────────────────────────────────────────────────────────────────────────
#  EXCEL
# ─────────────────────────────────────────────────────────────────────────────
def _fill(h): return PatternFill('solid', fgColor=h)
def _side(): return Side(border_style='thin', color='B8CCE4')
B_THIN = Border(left=_side(), right=_side(), top=_side(), bottom=_side())


def _cel(ws, r, c, v='', bold=False, sz=10, bg=None, fg=None, ha='left', wrap=False):
    x = ws.cell(row=r, column=c, value=v)
    x.font = Font(name='Arial', bold=bold, size=sz, color=fg or '000000')
    x.alignment = Alignment(horizontal=ha, vertical='center', wrap_text=wrap)
    if bg: x.fill = _fill(bg)
    x.border = B_THIN
    return x


def gerar_excel(achados, mes, ano, output_path):
    wb = openpyxl.Workbook(); ws = wb.active
    ws.title = 'Relatório de Integridade'
    ws.sheet_view.showGridLines = False
    widths = {'A': 8, 'B': 46, 'C': 30, 'D': 12, 'E': 18, 'F': 30,
              'G': 18, 'H': 30, 'I': 16, 'J': 55}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    ws.merge_cells('A1:J2')
    c = ws['A1']
    c.value = (f'GOVERNO DO DISTRITO FEDERAL\n'
               f'Relatório de Integridade Contábil e Orçamentária — '
               f'{MESES[mes]}/{ano}')
    c.font = Font(name='Arial', bold=True, size=13, color='2E5C8A')
    c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 40

    header = ['Regra', 'Título', 'Base Normativa', 'Status', 'Valor A',
              'Rótulo A', 'Valor B', 'Rótulo B', 'Diferença', 'Observação']
    H = 4
    for j, htxt in enumerate(header, start=1):
        _cel(ws, H, j, htxt, bold=True, bg='2E5C8A', fg='FFFFFF', ha='center')

    for i, a in enumerate(achados, start=H + 1):
        _, bg = COR_STATUS.get(a.status, ('000000', 'FFFFFF'))
        status_txt = STATUS_LABEL.get(a.status, a.status)
        _cel(ws, i, 1, a.regra, bg=bg, ha='center')
        _cel(ws, i, 2, a.titulo, bg=bg, wrap=True)
        _cel(ws, i, 3, a.base_normativa, bg=bg, wrap=True)
        _cel(ws, i, 4, status_txt, bold=True, bg=bg, ha='center')
        _cel(ws, i, 5, a.valor_a, bg=bg, ha='right')
        if a.valor_a is not None: ws.cell(row=i, column=5).number_format = '#,##0.00'
        _cel(ws, i, 6, a.rotulo_a, bg=bg, wrap=True)
        _cel(ws, i, 7, a.valor_b, bg=bg, ha='right')
        if a.valor_b is not None: ws.cell(row=i, column=7).number_format = '#,##0.00'
        _cel(ws, i, 8, a.rotulo_b, bg=bg, wrap=True)
        dif = a.diferenca
        _cel(ws, i, 9, dif, bg=bg, ha='right')
        if dif is not None: ws.cell(row=i, column=9).number_format = '#,##0.00'
        _cel(ws, i, 10, a.observacao, bg=bg, wrap=True)

    ws.freeze_panes = f'A{H+1}'
    wb.save(output_path)
    print(f"  Excel salvo: {output_path}")


# ─────────────────────────────────────────────────────────────────────────────
#  PDF
# ─────────────────────────────────────────────────────────────────────────────
def _brl(v):
    if v is None:
        return '—'
    return f'{v:,.2f}'.replace(',', '§').replace('.', ',').replace('§', '.')


def gerar_pdf(achados, mes, ano, output_path):
    if not REPORTLAB_OK:
        print("  AVISO: reportlab não instalado -- PDF não gerado.")
        return

    doc = SimpleDocTemplate(str(output_path), pagesize=A4,
                             leftMargin=12*mm, rightMargin=12*mm,
                             topMargin=12*mm, bottomMargin=12*mm)
    styles = getSampleStyleSheet()
    elems = []

    st_t = ParagraphStyle('t', parent=styles['Normal'], fontName='Helvetica-Bold',
                           fontSize=15, textColor=colors.HexColor('#2E5C8A'), leading=18)
    st_s = ParagraphStyle('s', parent=styles['Normal'], fontSize=9, textColor=colors.grey)
    elems.append(Paragraph('GOVERNO DO DISTRITO FEDERAL', st_s))
    elems.append(Paragraph(f'Relatório de Integridade Contábil e Orçamentária — '
                            f'{MESES[mes]}/{ano}', st_t))
    elems.append(Paragraph(f'Gerado em {datetime.now():%d/%m/%Y às %H:%M:%S} · '
                            f'Fonte: PDFs oficiais publicados (PSIAG550) — '
                            f'BF, BO, BP, DFC, DVP, DMPL', st_s))
    elems.append(Spacer(1, 6*mm))

    contagem = {}
    for a in achados:
        contagem[a.status] = contagem.get(a.status, 0) + 1
    ordem = ['OK', 'DIVERGENTE', 'ALERTA', 'INFORMATIVO', 'NAO_VERIF']
    painel_dados = [[str(contagem.get(s, 0)) for s in ordem],
                     [STATUS_LABEL.get(s, s) for s in ordem]]
    painel = Table([[str(len(achados))] + painel_dados[0],
                    ['Total'] + painel_dados[1]],
                   colWidths=[28*mm]*6)
    ts_painel = [('GRID', (0, 0), (-1, -1), 0.4, colors.white),
                 ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                 ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                 ('FONTSIZE', (0, 0), (-1, 0), 16),
                 ('FONTSIZE', (0, 1), (-1, 1), 8),
                 ('TOPPADDING', (0, 0), (-1, -1), 8),
                 ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
                 ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#2E5C8A')),
                 ('TEXTCOLOR', (0, 0), (0, -1), colors.white)]
    for i, s in enumerate(ordem, start=1):
        fg, bg = COR_STATUS[s]
        ts_painel.append(('BACKGROUND', (i, 0), (i, -1), _hx(bg)))
        ts_painel.append(('TEXTCOLOR', (i, 0), (i, 0), _hx(fg)))
    painel.setStyle(TableStyle(ts_painel))
    elems.append(painel)
    elems.append(Spacer(1, 6*mm))

    prioritarios = [a for a in achados if a.status in ('DIVERGENTE', 'ALERTA')]
    if prioritarios:
        st_h2 = ParagraphStyle('h2', parent=styles['Normal'], fontName='Helvetica-Bold',
                                fontSize=12, textColor=colors.HexColor('#2E5C8A'))
        elems.append(Paragraph('ACHADOS PRIORITÁRIOS (DIVERGENTE / ALERTA)', st_h2))
        elems.append(Spacer(1, 2*mm))
        st_r = ParagraphStyle('r', fontName='Helvetica-Bold', fontSize=9)
        st_n = ParagraphStyle('n', fontName='Helvetica', fontSize=8.5, leading=11)
        st_o = ParagraphStyle('o', fontName='Helvetica-Oblique', fontSize=8, leading=10,
                              textColor=colors.HexColor('#444444'))
        for a in prioritarios:
            fg, bg = COR_STATUS[a.status]
            linhas = [[Paragraph(f'Regra {a.regra} — {a.titulo}', st_r),
                      Paragraph(a.status, st_r)]]
            tab = Table(linhas, colWidths=[150*mm, 26*mm])
            tab.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, -1), _hx(bg)),
                ('TEXTCOLOR', (1, 0), (1, 0), _hx(fg)),
                ('ALIGN', (1, 0), (1, 0), 'CENTER'),
                ('TOPPADDING', (0, 0), (-1, -1), 4), ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ]))
            elems.append(tab)
            corpo = [[Paragraph(f'{a.rotulo_a or "Valor A"}: {_brl(a.valor_a)}', st_n),
                     Paragraph(f'{a.rotulo_b or "Valor B"}: {_brl(a.valor_b)}', st_n)]]
            if a.diferenca is not None:
                corpo.append([Paragraph(f'Diferença apurada: {_brl(a.diferenca)}', st_r), ''])
            tab2 = Table(corpo, colWidths=[88*mm, 88*mm])
            tab2.setStyle(TableStyle([('GRID', (0, 0), (-1, -1), 0.3, colors.lightgrey),
                                      ('TOPPADDING', (0, 0), (-1, -1), 3),
                                      ('BOTTOMPADDING', (0, 0), (-1, -1), 3)]))
            elems.append(tab2)
            if a.base_normativa or a.observacao:
                txt = (f'<b>Base normativa:</b> {a.base_normativa}. ' if a.base_normativa else '')
                txt += a.observacao
                elems.append(Paragraph(txt, st_o))
            elems.append(Spacer(1, 3*mm))
        elems.append(Spacer(1, 4*mm))

    st_h2 = ParagraphStyle('h2b', parent=styles['Normal'], fontName='Helvetica-Bold',
                            fontSize=12, textColor=colors.HexColor('#2E5C8A'))
    elems.append(Paragraph('DETALHAMENTO COMPLETO', st_h2))
    elems.append(Spacer(1, 2*mm))
    st_cell = ParagraphStyle('c', fontName='Helvetica', fontSize=7, leading=8.5)
    st_cell_b = ParagraphStyle('cb', fontName='Helvetica-Bold', fontSize=7, leading=8.5)
    st_val = ParagraphStyle('v', fontName='Helvetica', fontSize=7, alignment=TA_RIGHT, leading=8.5)
    dados = [[Paragraph('Regra', st_cell_b), Paragraph('Título', st_cell_b),
              Paragraph('Valor A', st_cell_b), Paragraph('Valor B', st_cell_b),
              Paragraph('Diferença', st_cell_b), Paragraph('Status', st_cell_b)]]
    for a in achados:
        dados.append([Paragraph(str(a.regra), st_cell), Paragraph(a.titulo, st_cell),
                      Paragraph(_brl(a.valor_a), st_val), Paragraph(_brl(a.valor_b), st_val),
                      Paragraph(_brl(a.diferenca), st_val),
                      Paragraph(STATUS_LABEL.get(a.status, a.status), st_cell_b)])
    tab = Table(dados, colWidths=[14*mm, 70*mm, 30*mm, 30*mm, 25*mm, 22*mm], repeatRows=1)
    ts = [('GRID', (0, 0), (-1, -1), 0.3, colors.grey),
          ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
          ('TOPPADDING', (0, 0), (-1, -1), 2), ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
          ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2E5C8A')),
          ('TEXTCOLOR', (0, 0), (-1, 0), colors.white)]
    for i, a in enumerate(achados, start=1):
        _, bg = COR_STATUS.get(a.status, ('000000', 'FFFFFF'))
        ts.append(('BACKGROUND', (0, i), (-1, i), _hx(bg)))
    tab.setStyle(TableStyle(ts))
    elems.append(tab)

    doc.build(elems)
    print(f"  PDF salvo:   {output_path}")


# ─────────────────────────────────────────────────────────────────────────────
#  PAINEL HTML — mesmo padrão visual/de histórico do auditoria_consolidada.py
#  e do gerar_balancete.py (seletor de execuções por data, snapshot em JSON).
# ─────────────────────────────────────────────────────────────────────────────
PAINEL_DIR = Path(__file__).parent / "painel"
OUTPUT_HTML = PAINEL_DIR / "painel_integridade_pdf_oficial.html"
DIR_HIST = PAINEL_DIR / "dados" / "integridade_pdf_oficial"
N_HISTORICO = 24

CSS_EXTRA = """
.c-gray{background:var(--s2);color:var(--t2);border-color:var(--bd)}
.kpi-row{grid-template-columns:repeat(5,1fr)}
@media (max-width:900px){.kpi-row{grid-template-columns:repeat(2,1fr)}}
"""


def salvar_historico(mes: int, ano: int, gerado_em: str, achados: list) -> None:
    """Grava um snapshot desta execução em painel/dados/integridade_pdf_oficial/,
    mesmo padrão do diag.py/auditoria_consolidada.py/gerar_balancete.py."""
    DIR_HIST.mkdir(parents=True, exist_ok=True)
    ts = gerado_em.replace("-", "").replace(":", "").replace("T", "_")[:15]
    caminho = DIR_HIST / f"{ano}-{mes:02d}_{ts}.json"
    doc = {
        "gerado_em": gerado_em,
        "mes": mes,
        "ano": ano,
        "achados": [
            {"regra": a.regra, "status": a.status, "titulo": a.titulo,
             "base_normativa": a.base_normativa,
             "valor_a": a.valor_a, "rotulo_a": a.rotulo_a,
             "valor_b": a.valor_b, "rotulo_b": a.rotulo_b,
             "diferenca": a.diferenca, "observacao": a.observacao}
            for a in achados
        ],
    }
    caminho.write_text(json.dumps(doc, ensure_ascii=False, indent=2), encoding="utf-8")


def carregar_historico(n: int = N_HISTORICO) -> list[dict]:
    """Últimas n execuções (mais recente primeiro), para o seletor de datas."""
    if not DIR_HIST.exists():
        return []
    arquivos = sorted(DIR_HIST.glob("*.json"), reverse=True)[:n]
    result = []
    for arq in arquivos:
        try:
            result.append(json.loads(arq.read_text(encoding="utf-8")))
        except Exception:
            pass
    return result


def _kpi(rotulo: str, valor: int, cor_token: str, val_id: str) -> str:
    return (f'<div class="kpi" style="--kpi-stripe:var({cor_token});--kpi-color:var({cor_token})">'
            f'<div class="kpi-lbl">{rotulo}</div><div class="kpi-val" id="{val_id}">{valor}</div></div>')


def _detalhe_txt(a) -> str:
    partes = []
    if a.get("rotulo_a") if isinstance(a, dict) else a.rotulo_a:
        ra = a["rotulo_a"] if isinstance(a, dict) else a.rotulo_a
        va = a["valor_a"] if isinstance(a, dict) else a.valor_a
        partes.append(f"{ra}: {fmt_brl(va)}" if va is not None else ra)
    rb = a["rotulo_b"] if isinstance(a, dict) else a.rotulo_b
    vb = a["valor_b"] if isinstance(a, dict) else a.valor_b
    if rb:
        partes.append(f"{rb}: {fmt_brl(vb)}" if vb is not None else rb)
    dif = a["diferenca"] if isinstance(a, dict) else a.diferenca
    if dif is not None:
        partes.append(f"Diferença: {fmt_brl(dif)}")
    obs = a["observacao"] if isinstance(a, dict) else a.observacao
    if obs:
        partes.append(obs)
    return "  |  ".join(partes)


def _linha_achado_html(a) -> str:
    status = a.status
    chip = CHIP_CLASS_RI.get(status, "c-inf")
    icone = ICONES_RI.get(status, "•")
    label = STATUS_LABEL.get(status, status)
    return (f'<div class="aud-item"><span class="chip {chip}">{icone} {label}</span>'
            f'<div class="aud-txt"><strong>{a.regra} — {a.titulo}</strong>'
            f'<span>{_detalhe_txt(a)}</span></div></div>')


def gerar_html(mes: int, ano: int, achados: list) -> str:
    agora_dt = datetime.now()
    agora = agora_dt.strftime("%d/%m/%Y %H:%M:%S")
    agora_iso = agora_dt.isoformat(timespec="seconds")

    cont = {"OK": 0, "DIVERGENTE": 0, "ALERTA": 0, "INFORMATIVO": 0, "NAO_VERIF": 0}
    for a in achados:
        cont[a.status] = cont.get(a.status, 0) + 1

    resultado_geral = (f"{cont['DIVERGENTE']} divergência(s), {cont['ALERTA']} alerta(s)"
                        if (cont["DIVERGENTE"] or cont["ALERTA"]) else "nenhuma divergência ou alerta")

    aud_html = "".join(_linha_achado_html(a) for a in achados)

    salvar_historico(mes, ano, agora_iso, achados)
    historico = carregar_historico()
    js_data = (
        f"const ICONES = {json.dumps(ICONES_RI, ensure_ascii=False)};\n"
        f"const CHIP_CLASS = {json.dumps(CHIP_CLASS_RI, ensure_ascii=False)};\n"
        f"const STATUS_LABEL = {json.dumps(STATUS_LABEL, ensure_ascii=False)};\n"
        f"const MESES = {json.dumps(MESES, ensure_ascii=False)};\n"
        f"var HISTORICO_FULL = {json.dumps(historico, ensure_ascii=False, indent=2)};"
    )

    return f"""<!DOCTYPE html>
<html lang="pt-BR">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>Integridade PDF Oficial</title>
<style>{_CSS_BASE}{CSS_EXTRA}</style>
</head>
<body>
<header class="hd">
  <div><div class="hd-title">Integridade Contábil — PDFs Oficiais (GDF)</div>
  <div class="hd-sub">{MESES[mes]}/{ano} · Fonte: PDFs publicados na pasta "13 - Demonstrativos Contábeis" (PSIAG550) — sem Oracle, sem xlsx interno</div></div>
  <div class="hd-sub">Gerado em {agora}</div>
</header>
<main class="wrap">
  <div class="run-wrap">
    <span class="run-lbl">Execução</span>
    <select id="run-select" onchange="loadRun(this.selectedIndex)"></select>
    <span class="run-note">seletor troca a auditoria pela execução escolhida</span>
  </div>

  <div class="kpi-row">
    {_kpi("Divergências", cont["DIVERGENTE"], "--err", "kpi-div")}
    {_kpi("Alertas", cont["ALERTA"], "--wrn", "kpi-alr")}
    {_kpi("OK", cont["OK"], "--ok", "kpi-ok")}
    {_kpi("Informativos", cont["INFORMATIVO"], "--inf", "kpi-inf")}
    {_kpi("Não Verificáveis", cont["NAO_VERIF"], "--t2", "kpi-nv")}
  </div>

  <details class="dem" open>
    <summary>Regras de Integridade <span id="resultado-geral">RESULTADO: {resultado_geral}</span></summary>
    <div class="dem-body"><div class="aud-list" id="aud-body">{aud_html}</div></div>
  </details>

  <footer>Integridade PDF Oficial — gerado automaticamente por relatorio_integridade_pdf_oficial.py (lê só os PDFs oficiais publicados, sem Oracle).</footer>
</main>
<script>
{js_data}

function fmtBRL(n){{
  if (n === null || n === undefined) return '';
  const neg = n < 0;
  let s = Math.abs(n).toLocaleString('pt-BR', {{minimumFractionDigits:2, maximumFractionDigits:2}});
  return (neg?'-':'') + s;
}}

function detalheTxt(a){{
  const partes=[];
  if(a.rotulo_a) partes.push(a.valor_a!=null ? `${{a.rotulo_a}}: ${{fmtBRL(a.valor_a)}}` : a.rotulo_a);
  if(a.rotulo_b) partes.push(a.valor_b!=null ? `${{a.rotulo_b}}: ${{fmtBRL(a.valor_b)}}` : a.rotulo_b);
  if(a.diferenca!=null) partes.push('Diferença: '+fmtBRL(a.diferenca));
  if(a.observacao) partes.push(a.observacao);
  return partes.join('  |  ');
}}

function renderRun(h){{
  const set=(id,v)=>{{const el=document.getElementById(id); if(el) el.textContent=v||0;}};
  const cont={{OK:0,DIVERGENTE:0,ALERTA:0,INFORMATIVO:0,NAO_VERIF:0}};
  (h.achados||[]).forEach(a=>{{cont[a.status]=(cont[a.status]||0)+1;}});
  set('kpi-div',cont.DIVERGENTE); set('kpi-alr',cont.ALERTA); set('kpi-ok',cont.OK);
  set('kpi-inf',cont.INFORMATIVO); set('kpi-nv',cont.NAO_VERIF);
  const rg=document.getElementById('resultado-geral');
  if(rg){{
    rg.textContent='RESULTADO: '+((cont.DIVERGENTE||cont.ALERTA)?`${{cont.DIVERGENTE}} divergência(s), ${{cont.ALERTA}} alerta(s)`:'nenhuma divergência ou alerta');
  }}
  const ab=document.getElementById('aud-body');
  if(ab){{
    ab.innerHTML=(h.achados||[]).map(a=>{{
      const chip=CHIP_CLASS[a.status]||'c-inf';
      const icone=ICONES[a.status]||'•';
      const label=STATUS_LABEL[a.status]||a.status;
      return `<div class="aud-item"><span class="chip ${{chip}}">${{icone}} ${{label}}</span>`+
             `<div class="aud-txt"><strong>${{a.regra}} — ${{a.titulo}}</strong><span>${{detalheTxt(a)}}</span></div></div>`;
    }}).join('');
  }}
}}

function buildSelector(){{
  const sel=document.getElementById('run-select');
  if(!sel||!HISTORICO_FULL.length) return;
  sel.innerHTML=HISTORICO_FULL.map((h,i)=>{{
    const nDiv=(h.achados||[]).filter(a=>a.status==='DIVERGENTE').length;
    let lbl='';
    try{{
      const d=new Date(h.gerado_em.replace('T',' '));
      lbl=`${{MESES[h.mes]||h.mes}}/${{h.ano}} — ${{d.toLocaleDateString('pt-BR')}} ${{d.toLocaleTimeString('pt-BR',{{hour:'2-digit',minute:'2-digit'}})}}`;
    }}catch(e){{lbl=h.gerado_em||String(i);}}
    if(nDiv) lbl+=' ⚠ '+nDiv+' divergência(s)';
    return `<option value="${{i}}">${{lbl}}</option>`;
  }}).join('');
}}

function loadRun(idx){{
  const h=HISTORICO_FULL[idx];
  if(!h) return;
  const sel=document.getElementById('run-select');
  if(sel) sel.selectedIndex=idx;
  renderRun(h);
}}

buildSelector();
</script>
</body>
</html>
"""


# ─────────────────────────────────────────────────────────────────────────────
#  MAIN
# ─────────────────────────────────────────────────────────────────────────────
def main():
    p = argparse.ArgumentParser(
        description='Relatório de Integridade Contábil e Orçamentária - GDF '
                    '(lê os PDFs oficiais da pasta 13, não precisa de Oracle)')
    p.add_argument('--mes', type=str, required=True,
                   help="1-12, ou 'ultimo' para o último mês com os 6 "
                        "demonstrativos já publicados na pasta 13")
    p.add_argument('--ano', type=int, required=True)
    p.add_argument('--formato', choices=['ambos', 'excel', 'pdf'], default='ambos')
    p.add_argument('--somente-se-mudou', action='store_true',
                   help='Só roda (e só grava arquivos) se a pasta 13 do ano '
                        'pedido mudou desde a última vez que este script '
                        'gerou o relatório com sucesso. Usado na rotina '
                        'noturna para não subir Excel/PDF novos sem motivo.')
    a = p.parse_args()

    if a.mes == 'ultimo':
        mes_resolvido = ultimo_mes_disponivel(a.ano)
        if mes_resolvido is None:
            print(f"ERRO: nenhum mês de {a.ano} tem os 6 demonstrativos "
                  f"publicados na pasta 13.")
            sys.exit(1)
        a.mes = mes_resolvido
        print(f"  --mes ultimo -> resolvido para {a.mes:02d}/{a.ano}")
    else:
        try:
            a.mes = int(a.mes)
        except ValueError:
            print("ERRO: --mes deve ser um número de 1 a 12, ou 'ultimo'."); sys.exit(1)

    if not 1 <= a.mes <= 12:
        print("ERRO: --mes deve estar entre 1 e 12."); sys.exit(1)

    if a.somente_se_mudou and not pasta13_mudou(a.ano):
        print(f"  Pasta 13 ({a.ano}) sem alterações desde a última rodada — "
              f"nada a gerar. Use sem --somente-se-mudou para forçar.")
        sys.exit(0)

    print(f"\n{'='*70}")
    print(f"  RELATÓRIO DE INTEGRIDADE CONTÁBIL E ORÇAMENTÁRIA")
    print(f"  {MESES[a.mes]}/{a.ano}  |  Fonte: PDFs oficiais (pasta 13)")
    print(f"{'='*70}")

    print("\n[1/3] Localizando e lendo os 6 PDFs oficiais...")
    try:
        d = carregar(a.mes, a.ano)
    except PdfNaoEncontrado as e:
        print(f"ERRO: {e}")
        sys.exit(1)
    for tipo, caminho in d.caminhos.items():
        print(f"  {tipo:<5} -> {caminho}")

    print("\n[2/3] Aplicando as regras de integridade...")
    achados = run_all_rules(d)
    from collections import Counter
    contagem = Counter(x.status for x in achados)
    print(f"  Total de verificações: {len(achados)}")
    for status, n in contagem.most_common():
        print(f"    {STATUS_LABEL.get(status, status)}: {n}")

    divergentes = [x for x in achados if x.status == 'DIVERGENTE']
    if divergentes:
        print(f"\n  {len(divergentes)} DIVERGÊNCIA(S):")
        for x in divergentes:
            print(f"    Regra {x.regra}: {x.titulo}")
            print(f"      {x.rotulo_a}: {x.valor_a:,.2f}  |  {x.rotulo_b}: {x.valor_b:,.2f}"
                  f"  |  dif: {x.diferenca:,.2f}")

    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    base = OUTPUT_DIR / f'Relatorio_Integridade_{a.ano}_{a.mes:02d}_{timestamp}'
    print("\n[3/3] Gerando arquivos...")
    if a.formato in ('ambos', 'excel'):
        gerar_excel(achados, a.mes, a.ano, base.with_suffix('.xlsx'))
    if a.formato in ('ambos', 'pdf'):
        gerar_pdf(achados, a.mes, a.ano, base.with_suffix('.pdf'))

    html = gerar_html(a.mes, a.ano, achados)
    OUTPUT_HTML.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_HTML.write_text(html, encoding='utf-8')
    print(f"  HTML salvo: {OUTPUT_HTML}")

    if a.somente_se_mudou:
        marcar_pasta13_processada(a.ano)

    print(f"\n  Concluído em {datetime.now():%H:%M:%S}\n")
    sys.exit(1 if divergentes else 0)


if __name__ == "__main__":
    main()
