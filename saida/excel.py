# -*- coding: utf-8 -*-
"""Gera um Excel de diagnóstico: uma aba, uma linha por controle."""
from __future__ import annotations
from datetime import datetime
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from controles import Achado
from saida.console import contar

# Paleta
COR = {
    "OK":     "D9F0DD",  # verde claro
    "ERRO":   "FAD7DA",  # vermelho claro
    "ALERTA": "FCE4D6",  # laranja claro
    "INFO":   "DEEAF1",  # azul claro
}
HDR_BG  = "2E5C8A"
HDR_FG  = "FFFFFF"

def _fill(h):  return PatternFill("solid", fgColor=h)
def _side(s="thin"): return Side(border_style=s, color="B8CCE4")
BORDA = Border(left=_side(), right=_side(), top=_side(), bottom=_side())


def _cel(ws, row, col, val="", bold=False, sz=9,
         bg=None, ha="left", color="000000"):
    c = ws.cell(row=row, column=col, value=val)
    c.font = Font(name="Arial", bold=bold, size=sz, color=color)
    c.alignment = Alignment(horizontal=ha, vertical="center")
    c.border = BORDA
    if bg: c.fill = _fill(bg)
    return c


def gerar(achados: list[Achado], mes: int, ano: int,
          output_path: Path,
          escopo: str = "Consolidado",
          emitido_em: datetime | None = None):
    """
    escopo: "Consolidado" ou o código da UG. Sem isso, dois arquivos de
        competências iguais e escopos diferentes ficam indistinguíveis para
        quem recebe a planilha solta.
    emitido_em: carimbo do INÍCIO da execução. Se omitido, usa a hora da
        gravação — que difere do carimbo do nome do arquivo pelo tempo que o
        diagnóstico levou (14 minutos, na configuração atual).
    """
    nomes = {1:"Janeiro",2:"Fevereiro",3:"Março",4:"Abril",5:"Maio",
             6:"Junho",7:"Julho",8:"Agosto",9:"Setembro",10:"Outubro",
             11:"Novembro",12:"Dezembro"}
    quando = emitido_em or datetime.now()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Diagnóstico"
    ws.sheet_view.showGridLines = False

    # Larguras
    for col, w in zip("ABCDEF", [8, 12, 10, 48, 60, 18]):
        ws.column_dimensions[col].width = w

    # Cabeçalho do relatório
    ws.merge_cells("A1:F1")
    c = ws["A1"]
    c.value = (f"DIAGNÓSTICO CONTÁBIL GDF  ·  "
               f"{nomes.get(mes,mes)}/{ano}  ·  "
               f"{escopo}  ·  "
               f"Emitido em {quando:%d/%m/%Y %H:%M}")
    c.font = Font(name="Arial", bold=True, size=12, color=HDR_FG)
    c.fill = _fill(HDR_BG)
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 22

    # Cabeçalho das colunas
    H = 3
    for j, txt in enumerate(["Status","Módulo","Código",
                              "Controle","Detalhe","Diferença (R$)"], 1):
        _cel(ws, H, j, txt, bold=True, bg=HDR_BG, ha="center", color=HDR_FG)
    ws.row_dimensions[H].height = 18
    ws.freeze_panes = f"A{H+1}"

    # Linhas
    r = H + 1
    for a in achados:
        bg = COR.get(a.status, "FFFFFF")
        _cel(ws, r, 1, a.status,  bold=True, bg=bg, ha="center")
        _cel(ws, r, 2, a.modulo,  bg=bg, ha="center")
        _cel(ws, r, 3, a.codigo,  bg=bg, ha="center")
        _cel(ws, r, 4, a.titulo,  bg=bg)
        _cel(ws, r, 5, (a.detalhe or "")[:500], bg=bg, sz=8)
        if a.valor is not None:
            try:
                c = ws.cell(row=r, column=6, value=float(a.valor))
                c.font  = Font(name="Arial", bold=True, size=9)
                c.fill  = _fill(bg)
                c.border = BORDA
                c.number_format = "#,##0.00;(#,##0.00)"
                c.alignment = Alignment(horizontal="right", vertical="center")
            except (ValueError, TypeError):
                _cel(ws, r, 6, str(a.valor)[:20], bg=bg, ha="right")
        else:
            _cel(ws, r, 6, "-", bg=bg, ha="center")
        ws.row_dimensions[r].height = 15
        r += 1

    # Rodapé de contagem — mesma função do console, para as duas saídas nunca
    # divergirem. Inclui INFO: antes o total dizia 29 controles e a quebra
    # somava 27, porque os informativos ficavam de fora.
    r += 1
    n = contar(achados)
    ws.merge_cells(f"A{r}:F{r}")
    txt = (f"Total: {n['TOTAL']} controles  ·  "
           f"{n['ERRO']} erro(s)  ·  {n['ALERTA']} alerta(s)  ·  "
           f"{n['OK']} OK  ·  {n['INFO']} informativo(s)")
    _cel(ws, r, 1, txt, bold=True, bg=HDR_BG, ha="left", color=HDR_FG)
    ws.row_dimensions[r].height = 16

    # AutoFiltro
    ws.auto_filter.ref = f"A{H}:F{r-2}"

    wb.save(output_path)
    return output_path
